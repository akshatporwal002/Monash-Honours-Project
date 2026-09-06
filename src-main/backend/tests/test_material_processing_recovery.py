from __future__ import annotations

import asyncio
import io
import subprocess
import sys
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime, timedelta
from pathlib import Path
from threading import Barrier

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import func, select

from app.api.routes.materials import get_actor_id, get_course_access_policy, get_material_storage
from app.core.config import Settings
from app.db.base import Base
from app.db.session import create_db_engine, create_session_factory, get_db_session
from app.main import create_app
from app.models import LearningMaterial, MaterialChunk, MaterialIndexStatus
from app.models.source_history import SourcePassage, SourceRevision
from app.services.material_indexing import OfflineMaterialProcessor
from app.services.rag.contracts import ExtractedBlock, ExtractedDocument
from app.services.rag.errors import MaterialAlreadyProcessingError, RagError
from app.services.rag.fakes import (
    AllowAllCourseAccessPolicy,
    DeterministicEmbeddingProvider,
    InMemoryVectorStore,
    StaticDocumentExtractor,
)
from app.services.rag.ingestion import MaterialProcessor
from app.services.rag.processing_claims import (
    LostMaterialClaim,
    MaterialProcessingClaims,
    MaterialRetryRequired,
)
from app.services.rag.processing_recovery import MaterialRecoveryWorker
from app.services.rag.storage import LocalFileStorage

NOW = datetime(2026, 9, 7, 4, 0, tzinfo=UTC)


@pytest.fixture
def processing_context(tmp_path: Path):
    database = tmp_path / "processing.db"
    engine = create_db_engine(f"sqlite:///{database.as_posix()}")
    Base.metadata.create_all(engine)
    factory = create_session_factory(engine)
    config = Settings(
        rag_upload_dir=str(tmp_path / "uploads"),
        material_processing_lease_seconds=5,
        material_processing_retry_seconds=1,
    )
    storage = LocalFileStorage(config.rag_upload_dir, config.rag_max_file_bytes)
    stream = io.BytesIO()
    document = Document()
    document.add_heading("Hadamard", level=1)
    document.add_paragraph("A Hadamard gate prepares a superposition from the zero state.")
    document.save(stream)
    staged = storage.stage_upload("lesson.docx", io.BytesIO(stream.getvalue()))
    with factory() as session:
        material = LearningMaterial(
            course_id="recovery-course",
            original_filename="lesson.docx",
            mime_type=staged.mime_type,
            content_hash=staged.content_hash,
        )
        session.add(material)
        session.flush()
        material.storage_key = storage.commit(staged, material.id)
        session.commit()
        material_id = material.id
    clock = [NOW]
    yield factory, storage, material_id, config, clock, database
    engine.dispose()


def _semantic(session, storage, config, clock, body, embedding=None, vectors=None):
    document = ExtractedDocument(
        (ExtractedBlock(0, body, "Heading", "Slide 2", "paragraph"),), None, len(body)
    )
    return MaterialProcessor(
        session,
        storage,
        {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document": StaticDocumentExtractor(
                document
            )
        },
        embedding or DeterministicEmbeddingProvider(),
        vectors or InMemoryVectorStore(),
        now=lambda: clock[0],
        configured_settings=config,
    )


def test_two_claimants_have_one_winner(processing_context):
    factory, _, material_id, config, clock, _ = processing_context
    ready = Barrier(2)

    def claim():
        with factory() as session:
            material = session.get(LearningMaterial, material_id)
            ready.wait()
            try:
                return MaterialProcessingClaims(
                    session, now=lambda: clock[0], configured_settings=config
                ).claim(material, backend="offline")
            except MaterialAlreadyProcessingError:
                return None

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda _: claim(), range(2)))
    assert sum(result is not None for result in results) == 1
    with factory() as session:
        material = session.get(LearningMaterial, material_id)
        assert material.processing_attempts == 1
        assert material.indexing_status is MaterialIndexStatus.PROCESSING


def test_worker_recovers_a_claim_from_a_terminated_process(processing_context):
    factory, storage, material_id, config, clock, database = processing_context
    # A separate interpreter exits without cleanup after its claim is committed.
    child = subprocess.run(
        [
            sys.executable,
            "-c",
            """
import os, sys
from datetime import datetime
from app.core.config import Settings
from app.db.session import create_db_engine, create_session_factory
from app.models import LearningMaterial
from app.services.rag.processing_claims import MaterialProcessingClaims
factory = create_session_factory(create_db_engine('sqlite:///' + sys.argv[1]))
with factory() as session:
    claim = MaterialProcessingClaims(session, now=lambda: datetime.fromisoformat(sys.argv[3]), configured_settings=Settings(material_processing_lease_seconds=5)).claim(session.get(LearningMaterial, sys.argv[2]), backend='offline')
    os._exit(17)
""",
            database.as_posix(),
            material_id,
            NOW.isoformat(),
        ],
        cwd=Path(__file__).resolve().parents[1],
        capture_output=True,
        timeout=30,
    )
    assert child.returncode == 17
    with factory() as session:
        material = session.get(LearningMaterial, material_id)
        old_token = material.processing_token
        original_key = material.storage_key
        assert old_token
        assert material.indexing_status is MaterialIndexStatus.PROCESSING
    worker = MaterialRecoveryWorker(factory, now=lambda: clock[0], configured_settings=config)
    assert asyncio.run(worker.run_once()) is False
    clock[0] += timedelta(seconds=6)
    assert asyncio.run(worker.run_once()) is True
    assert asyncio.run(worker.run_once()) is False
    with factory() as session:
        material = session.get(LearningMaterial, material_id)
        assert material.indexing_status is MaterialIndexStatus.INDEXED
        assert material.processing_token is None
        assert material.processing_attempts == 2
        assert session.scalar(select(func.count()).select_from(SourceRevision)) == 1
        assert "Hadamard" in session.scalar(select(SourcePassage.chunk_text))
        assert (storage.upload_dir / original_key).is_file()


def test_late_semantic_completion_cannot_replace_the_recovered_revision(processing_context):
    factory, storage, material_id, config, clock, _ = processing_context
    vectors = InMemoryVectorStore()

    class ReclaimedWhileEmbedding(DeterministicEmbeddingProvider):
        def embed_documents(self, texts):
            clock[0] += timedelta(seconds=6)
            with factory() as winner_session:
                winner = winner_session.get(LearningMaterial, material_id)
                _semantic(
                    winner_session,
                    storage,
                    config,
                    clock,
                    "Recovered evidence about the quantum measurement process.",
                    vectors=vectors,
                ).process(winner, recover=True)
            return super().embed_documents(texts)

    with factory() as session:
        material = session.get(LearningMaterial, material_id)
        with pytest.raises(LostMaterialClaim):
            _semantic(
                session,
                storage,
                config,
                clock,
                "Late stale evidence that must never be published.",
                embedding=ReclaimedWhileEmbedding(),
                vectors=vectors,
            ).process(material)
    with factory() as session:
        material = session.get(LearningMaterial, material_id)
        assert material.indexing_status is MaterialIndexStatus.INDEXED
        assert material.processing_attempts == 2
        assert session.scalar(select(func.count()).select_from(SourceRevision)) == 1
        assert (
            session.scalar(select(SourcePassage.chunk_text))
            == "Recovered evidence about the quantum measurement process."
        )
        assert (
            session.scalar(select(MaterialChunk.chunk_text))
            == "Recovered evidence about the quantum measurement process."
        )


def test_failed_processing_retries_when_due_without_exposing_provider_details(processing_context):
    factory, storage, material_id, config, clock, _ = processing_context

    class BrokenEmbedding(DeterministicEmbeddingProvider):
        def embed_documents(self, texts):
            raise RuntimeError("private provider response with credentials")

    with factory() as session:
        material = session.get(LearningMaterial, material_id)
        with pytest.raises(RuntimeError):
            _semantic(
                session,
                storage,
                config,
                clock,
                "Evidence for the quantum source recovery test.",
                embedding=BrokenEmbedding(),
            ).process(material)
        assert "private" not in material.extraction_error
        assert material.processing_retry_at is not None
    worker = MaterialRecoveryWorker(
        factory,
        now=lambda: clock[0],
        configured_settings=config,
        processor_factory=lambda session, backend: _semantic(
            session, storage, config, clock, "Evidence for the quantum source recovery test."
        ),
    )
    assert asyncio.run(worker.run_once()) is False
    clock[0] += timedelta(seconds=1)
    assert asyncio.run(worker.run_once()) is True
    with factory() as session:
        material = session.get(LearningMaterial, material_id)
        assert material.indexing_status is MaterialIndexStatus.INDEXED
        assert material.processing_attempts == 2
        assert material.processing_retry_at is None
        assert material.extraction_error is None


def test_final_interrupted_claim_stops_until_an_explicit_fresh_run(processing_context):
    factory, storage, material_id, config, clock, _ = processing_context
    for _ in range(3):
        with factory() as session:
            MaterialProcessingClaims(
                session, now=lambda: clock[0], configured_settings=config
            ).claim(session.get(LearningMaterial, material_id), backend="offline", recover=True)
        clock[0] += timedelta(seconds=6)
    worker = MaterialRecoveryWorker(factory, now=lambda: clock[0], configured_settings=config)
    assert asyncio.run(worker.run_once()) is True
    assert asyncio.run(worker.run_once()) is False
    with factory() as session:
        material = session.get(LearningMaterial, material_id)
        assert material.indexing_status is MaterialIndexStatus.FAILED
        assert material.error_code == "material_processing_attempts_exhausted"
        assert material.processing_attempts == 3
        processor = OfflineMaterialProcessor(
            session, storage, now=lambda: clock[0], configured_settings=config
        )
        with pytest.raises(MaterialRetryRequired):
            processor.process(material)
        processor.process(material, force=True)
        assert material.indexing_status is MaterialIndexStatus.INDEXED
        assert material.processing_attempts == 1
        assert material.processing_revision == 1


def test_missing_semantic_adapter_does_not_silently_switch_to_offline(processing_context):
    factory, _, material_id, config, clock, _ = processing_context
    with factory() as session:
        MaterialProcessingClaims(session, now=lambda: clock[0], configured_settings=config).claim(
            session.get(LearningMaterial, material_id), backend="semantic"
        )
    clock[0] += timedelta(seconds=6)
    worker = MaterialRecoveryWorker(factory, now=lambda: clock[0], configured_settings=config)
    assert asyncio.run(worker.run_once()) is True
    with factory() as session:
        material = session.get(LearningMaterial, material_id)
        assert material.processing_backend == "semantic"
        assert material.error_code == "material_processor_unavailable"
        assert session.scalar(select(func.count()).select_from(SourceRevision)) == 0


@pytest.mark.parametrize(
    "state",
    [MaterialIndexStatus.PENDING, MaterialIndexStatus.PROCESSING, MaterialIndexStatus.FAILED],
)
def test_worker_recovers_legacy_saved_uploads_without_claim_metadata(processing_context, state):
    factory, _, material_id, config, clock, _ = processing_context
    with factory() as session:
        material = session.get(LearningMaterial, material_id)
        material.indexing_status = state
        session.commit()
    worker = MaterialRecoveryWorker(factory, now=lambda: clock[0], configured_settings=config)
    assert asyncio.run(worker.run_once()) is True
    with factory() as session:
        material = session.get(LearningMaterial, material_id)
        assert material.indexing_status is MaterialIndexStatus.INDEXED
        assert material.processing_attempts == 1


def test_expiry_during_index_publication_rolls_back_the_partial_revision(processing_context):
    factory, storage, material_id, config, clock, _ = processing_context

    class SlowVectorStore(InMemoryVectorStore):
        def upsert(self, records):
            super().upsert(records)
            clock[0] += timedelta(seconds=6)

    with factory() as session:
        material = session.get(LearningMaterial, material_id)
        with pytest.raises(LostMaterialClaim):
            _semantic(
                session,
                storage,
                config,
                clock,
                "This publication exceeds its processing lease deadline.",
                vectors=SlowVectorStore(),
            ).process(material)
        assert session.scalar(select(func.count()).select_from(SourceRevision)) == 0
        assert session.scalar(select(func.count()).select_from(MaterialChunk)) == 0
    with factory() as session:
        material = session.get(LearningMaterial, material_id)
        _semantic(
            session,
            storage,
            config,
            clock,
            "The recovered publication completes within its deadline.",
        ).process(material, recover=True)
        assert session.scalar(select(func.count()).select_from(SourceRevision)) == 1


def test_broken_recovery_factory_records_a_bounded_failure(processing_context):
    factory, _, material_id, config, clock, _ = processing_context

    def broken_factory(session, backend):
        raise RuntimeError("private adapter startup failure")

    worker = MaterialRecoveryWorker(
        factory, now=lambda: clock[0], configured_settings=config, processor_factory=broken_factory
    )
    assert asyncio.run(worker.run_once()) is True
    assert asyncio.run(worker.run_once()) is False
    with factory() as session:
        material = session.get(LearningMaterial, material_id)
        assert material.processing_attempts == 1
        assert material.error_code == "material_processor_unavailable"
        assert "private" not in material.extraction_error


def test_material_api_exposes_safe_status_and_allows_an_explicit_fresh_run(processing_context):
    factory, storage, material_id, _, _, _ = processing_context
    with factory() as session:
        material = session.get(LearningMaterial, material_id)
        material.indexing_status = MaterialIndexStatus.FAILED
        material.processing_attempts = 3
        material.error_code = "material_processing_attempts_exhausted"
        material.extraction_error = "Processing stopped. Review the saved upload and retry."
        session.commit()

    def database():
        with factory() as session:
            yield session

    app = create_app()
    app.dependency_overrides[get_db_session] = database
    app.dependency_overrides[get_actor_id] = lambda: "course-owner"
    app.dependency_overrides[get_course_access_policy] = AllowAllCourseAccessPolicy
    app.dependency_overrides[get_material_storage] = lambda: storage
    base = f"/api/v1/courses/recovery-course/materials/{material_id}"
    with TestClient(app) as client:
        status = client.get(base).json()
        assert status["processing_attempts"] == 3
        assert status["processing_retry_at"] is None
        assert status["error_code"] == "material_processing_attempts_exhausted"
        assert "processing_token" not in status
        assert client.post(f"{base}/process").status_code == 409
        retried = client.post(f"{base}/process?force=true")
        assert retried.status_code == 200, retried.text
        assert retried.json()["material"]["processing_attempts"] == 1
        assert retried.json()["material"]["indexing_status"] == "indexed"
        assert retried.json()["material"]["extraction_error"] is None
        with factory() as session:
            claims = MaterialProcessingClaims(session, now=lambda: NOW)
            claims.claim(session.get(LearningMaterial, material_id), backend="offline", force=True)
        deadline = datetime.fromisoformat(client.get(base).json()["processing_lease_expires_at"])
        assert deadline.tzinfo is not None
        assert deadline > NOW

        assert (
            client.post(
                f"/api/v1/courses/another-course/materials/{material_id}/process?force=true"
            ).status_code
            == 404
        )


def test_manual_retry_cannot_silently_replace_the_saved_processing_backend(processing_context):
    factory, storage, material_id, config, clock, _ = processing_context
    with factory() as session:
        claims = MaterialProcessingClaims(session, now=lambda: clock[0], configured_settings=config)
        material = session.get(LearningMaterial, material_id)
        claim = claims.claim(material, backend="semantic")
        claims.fail(claim, RuntimeError("provider unavailable"))
        with pytest.raises(RagError, match="saved processing adapter"):
            OfflineMaterialProcessor(
                session, storage, now=lambda: clock[0], configured_settings=config
            ).process(material, force=True)
        session.refresh(material)
        assert material.processing_backend == "semantic"
        assert material.processing_attempts == 1
        assert material.indexing_status is MaterialIndexStatus.FAILED
